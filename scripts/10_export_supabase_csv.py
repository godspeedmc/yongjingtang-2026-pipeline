from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

STEMS = "甲乙丙丁戊己庚辛壬癸"
BRANCHES = "子丑寅卯辰巳午未申酉戌亥"
VALID_PILLARS = {STEMS[i % 10] + BRANCHES[i % 12] for i in range(60)}
WEALTH = {"甲":"東北","乙":"東北","丙":"西南","丁":"西南","戊":"正北","己":"正北","庚":"正東","辛":"正東","壬":"正南","癸":"正南"}
JOY = {"甲":"東北","乙":"西北","丙":"西南","丁":"正南","戊":"東南","己":"東北","庚":"西北","辛":"西南","壬":"正南","癸":"東南"}
CLASH = {"子":"馬","丑":"羊","寅":"猴","卯":"雞","辰":"狗","巳":"豬","午":"鼠","未":"牛","申":"虎","酉":"兔","戌":"龍","亥":"蛇"}
SHA = {"子":"南","丑":"東","寅":"北","卯":"西","辰":"南","巳":"東","午":"北","未":"西","申":"南","酉":"東","戌":"北","亥":"西"}

OUTPUT_FIELDS = [
    "date","edition","source_pdf_page","printed_page","lunar_date","day_pillar",
    "wealth_direction","joy_direction","clash_zodiac","sha_direction",
    "lunar_date_source","day_pillar_source","wealth_direction_source",
    "joy_direction_source","clash_zodiac_source","sha_direction_source",
    "verification_status","verification_notes","reviewed_by_human","verified_at",
    "extraction_method","extraction_confidence",
]


def parse_args():
    p = argparse.ArgumentParser(description="Export human-reviewed workbook rows to a Supabase-ready CSV.")
    p.add_argument("--workbook", default="review/yongjingtang_review_workbook.xlsx")
    p.add_argument("--output", default="verified/yongjingtang_2026_hybrid_publishable.csv")
    p.add_argument("--report", default="reports/supabase_export_report.json")
    p.add_argument("--edition", default="香港永經堂通勝 2026")
    return p.parse_args()


def clean(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def main():
    args = parse_args()
    root = Path.cwd()
    workbook = root / args.workbook
    output = root / args.output
    report_path = root / args.report
    output.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)

    if not workbook.exists():
        raise FileNotFoundError(f"Review workbook not found: {workbook}")

    wb = load_workbook(workbook, data_only=True)
    if "Review" not in wb.sheetnames:
        raise ValueError("Workbook does not contain a Review sheet.")
    ws = wb["Review"]
    headers = [clean(cell.value) for cell in ws[1]]
    index = {name: position for position, name in enumerate(headers)}
    required = {"source_pdf_page","column_index","date","printed_page","lunar_date","day_pillar","verification_status","review_complete"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError("Missing workbook columns: " + ", ".join(missing))

    reviewed = []
    rejected = []
    for row_number, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row = {header: clean(cells[position]) for header, position in index.items()}
        if row.get("review_complete", "").lower() != "yes" or row.get("verification_status") != "human_verified":
            continue

        errors = []
        pillar = row.get("day_pillar", "")
        try:
            datetime.strptime(row.get("date", ""), "%Y-%m-%d")
        except ValueError:
            errors.append("invalid date")
        if pillar not in VALID_PILLARS:
            errors.append("invalid day pillar")
        if not row.get("printed_page"):
            errors.append("missing printed page")
        if not row.get("lunar_date"):
            errors.append("missing lunar date")

        expected = {
            "wealth_direction": WEALTH.get(pillar[:1], ""),
            "joy_direction": JOY.get(pillar[:1], ""),
            "clash_zodiac": CLASH.get(pillar[1:2], ""),
            "sha_direction": SHA.get(pillar[1:2], ""),
        }
        for field, value in expected.items():
            entered = row.get(field, "")
            if entered and entered != value:
                errors.append(f"{field} conflicts with generic rule")
            row[field] = value

        if errors:
            rejected.append({"row": row_number, "date": row.get("date"), "errors": errors})
            continue

        reviewed.append({
            "date": row["date"],
            "edition": args.edition,
            "source_pdf_page": row["source_pdf_page"],
            "printed_page": row["printed_page"],
            "lunar_date": row["lunar_date"],
            "day_pillar": pillar,
            "wealth_direction": row["wealth_direction"],
            "joy_direction": row["joy_direction"],
            "clash_zodiac": row["clash_zodiac"],
            "sha_direction": row["sha_direction"],
            "lunar_date_source": "yong_jing_tang",
            "day_pillar_source": "yong_jing_tang",
            "wealth_direction_source": "generic_rule",
            "joy_direction_source": "generic_rule",
            "clash_zodiac_source": "generic_rule",
            "sha_direction_source": "generic_rule",
            "verification_status": "human_verified",
            "verification_notes": row.get("verification_notes", ""),
            "reviewed_by_human": True,
            "verified_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "extraction_method": "human_reviewed_publication_fields+generic_rule_derived_fields",
            "extraction_confidence": 1.0,
        })

    counts = Counter(item["date"] for item in reviewed)
    duplicates = sorted(date for date, count in counts.items() if count > 1)
    if duplicates:
        rejected.append({"row": None, "date": None, "errors": ["duplicate dates: " + ", ".join(duplicates)]})

    report = {
        "workbook": str(workbook),
        "publishable_records": len(reviewed),
        "rejected": rejected,
        "duplicate_dates": duplicates,
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    if rejected:
        print("Export blocked. Review the report:")
        print(report_path)
        raise SystemExit(1)

    reviewed.sort(key=lambda item: item["date"])
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_FIELDS)
        writer.writeheader()
        writer.writerows(reviewed)

    print(f"Supabase-ready CSV: {output}")
    print(f"Publishable records: {len(reviewed)}")
    print(f"Export report: {report_path}")


if __name__ == "__main__":
    main()
