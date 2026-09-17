from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


FIELDS = [
    "source_pdf_page", "column_index", "date", "printed_page",
    "lunar_date", "day_pillar", "wealth_direction", "joy_direction",
    "clash_zodiac", "sha_direction", "lunar_date_source",
    "day_pillar_source", "wealth_direction_source", "joy_direction_source",
    "clash_zodiac_source", "sha_direction_source", "verification_status",
    "verification_notes", "source_crop", "pillar_ocr", "pillar_votes",
    "review_complete",
]


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Create an Excel review workbook from candidate JSONL and existing reviewed CSVs."
    )
    p.add_argument("--candidates", default="extraction/yongjingtang_2026_candidates.jsonl")
    p.add_argument("--pillar-summary", default="ocr-results/pillars/pillar_ocr_summary.csv")
    p.add_argument("--review-dir", default="review")
    p.add_argument("--output", default="review/yongjingtang_review_workbook.xlsx")
    p.add_argument("--pages", default="", help="Optional comma-separated PDF pages, for example 31,32.")
    return p.parse_args()


def resolve(root: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else root / path


def as_int(value: Any, default: int = 0) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def text(value: Any) -> str:
    return str(value or "").strip()


def first(record: dict[str, Any], key: str) -> str:
    value = record.get(key, [])
    return text(value[0]) if isinstance(value, list) and value else ""


def load_candidates(path: Path, requested: set[int]) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"Candidate JSONL not found: {path}")
    rows = []
    with path.open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            line = line.strip()
            if not line:
                continue
            try:
                record = json.loads(line)
            except json.JSONDecodeError as error:
                raise ValueError(f"Invalid JSONL line {line_number}: {error}") from error
            page = as_int(record.get("source_pdf_page"))
            if requested and page not in requested:
                continue
            rows.append(record)
    return rows


def load_pillars(path: Path) -> dict[tuple[int, int], dict[str, str]]:
    if not path.exists():
        return {}
    result = {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            result[(as_int(row.get("source_pdf_page")), as_int(row.get("column_index")))] = row
    return result


def load_existing_reviews(directory: Path) -> dict[tuple[int, int], dict[str, str]]:
    result = {}
    if not directory.exists():
        return result
    for path in sorted(directory.glob("ground_truth_page_*.csv")):
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            for row in csv.DictReader(handle):
                key = (as_int(row.get("source_pdf_page")), as_int(row.get("column_index")))
                if key[0] and key[1]:
                    result[key] = row
    return result


def style_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {
        "A": 16, "B": 12, "C": 14, "D": 14, "E": 16, "F": 14,
        "G": 18, "H": 16, "I": 15, "J": 14, "K": 19, "L": 19,
        "M": 24, "N": 22, "O": 20, "P": 20, "Q": 20, "R": 55,
        "S": 46, "T": 14, "U": 12, "V": 16,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def main() -> None:
    args = parse_args()
    root = Path.cwd()
    requested = {int(v.strip()) for v in args.pages.split(",") if v.strip()}
    candidates = load_candidates(resolve(root, args.candidates), requested)
    if not candidates:
        raise SystemExit("No candidate rows selected. Run Script 06 first or check --pages.")

    pillars = load_pillars(resolve(root, args.pillar_summary))
    existing = load_existing_reviews(resolve(root, args.review_dir))
    output = resolve(root, args.output)
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append(FIELDS)

    for candidate in sorted(candidates, key=lambda r: (as_int(r.get("source_pdf_page")), as_int(r.get("column_index")))):
        page = as_int(candidate.get("source_pdf_page"))
        column = as_int(candidate.get("column_index"))
        key = (page, column)
        reviewed = existing.get(key, {})
        pillar = pillars.get(key, {})
        crop = text(candidate.get("source_crop"))

        values = {
            "source_pdf_page": page,
            "column_index": column,
            "date": text(reviewed.get("date")) or text(candidate.get("date_candidate")),
            "printed_page": text(reviewed.get("printed_page")),
            "lunar_date": text(reviewed.get("lunar_date")) or first(candidate, "lunar_date_candidates"),
            "day_pillar": text(reviewed.get("day_pillar")) or text(pillar.get("day_pillar_top")) or first(candidate, "day_pillar_candidates"),
            "wealth_direction": text(reviewed.get("wealth_direction")),
            "joy_direction": text(reviewed.get("joy_direction")),
            "clash_zodiac": text(reviewed.get("clash_zodiac")),
            "sha_direction": text(reviewed.get("sha_direction")),
            "lunar_date_source": text(reviewed.get("lunar_date_source")) or "yong_jing_tang",
            "day_pillar_source": text(reviewed.get("day_pillar_source")) or "yong_jing_tang",
            "wealth_direction_source": text(reviewed.get("wealth_direction_source")) or "generic_rule",
            "joy_direction_source": text(reviewed.get("joy_direction_source")) or "generic_rule",
            "clash_zodiac_source": text(reviewed.get("clash_zodiac_source")) or "generic_rule",
            "sha_direction_source": text(reviewed.get("sha_direction_source")) or "generic_rule",
            "verification_status": text(reviewed.get("verification_status")) or "needs_review",
            "verification_notes": text(reviewed.get("verification_notes")),
            "source_crop": crop,
            "pillar_ocr": text(pillar.get("day_pillar_top")),
            "pillar_votes": as_int(pillar.get("day_pillar_votes")),
            "review_complete": "yes" if text(reviewed.get("verification_status")) == "human_verified" else "no",
        }
        ws.append([values[field] for field in FIELDS])
        row_number = ws.max_row
        crop_cell = ws.cell(row=row_number, column=FIELDS.index("source_crop") + 1)
        crop_path = resolve(root, crop) if crop else None
        if crop_path and crop_path.exists():
            crop_cell.hyperlink = crop_path.as_uri()
            crop_cell.style = "Hyperlink"

    style_sheet(ws)

    status_validation = DataValidation(
        type="list",
        formula1='"needs_review,human_verified,unclear_scan,source_conflict"',
        allow_blank=False,
    )
    complete_validation = DataValidation(type="list", formula1='"yes,no"', allow_blank=False)
    ws.add_data_validation(status_validation)
    ws.add_data_validation(complete_validation)
    status_col = get_column_letter(FIELDS.index("verification_status") + 1)
    complete_col = get_column_letter(FIELDS.index("review_complete") + 1)
    status_validation.add(f"{status_col}2:{status_col}{ws.max_row}")
    complete_validation.add(f"{complete_col}2:{complete_col}{ws.max_row}")

    instructions = wb.create_sheet("Instructions")
    instructions.append(["Review workflow"])
    instructions.append(["1. Open each source_crop hyperlink and verify the date, printed page, lunar date, and day pillar."])
    instructions.append(["2. Keep publication sources as yong_jing_tang and derived direction/clash fields as generic_rule."])
    instructions.append(["3. Set verification_status to human_verified and review_complete to yes only after visual review."])
    instructions.append(["4. Save the workbook, then export reviewed rows using Script 10."])
    instructions.column_dimensions["A"].width = 120
    instructions["A1"].font = Font(bold=True, color="FFFFFF")
    instructions["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    for row in instructions.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(output)
    print(f"Review workbook: {output}")
    print(f"Review rows: {ws.max_row - 1}")


if __name__ == "__main__":
    main()
